import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def crear_documento_word():
    doc = Document()

    # Configuración de márgenes normales (2.5 cm = ~0.98 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores institucionales
    COLOR_TITULO = RGBColor(15, 23, 42)      # Slate 900
    COLOR_PRIMARIO = RGBColor(29, 185, 84)   # Spotify Green
    COLOR_SUBTITULO = RGBColor(71, 85, 105)  # Slate 600
    COLOR_TEXTO = RGBColor(30, 41, 59)       # Slate 800
    COLOR_HEADER_BG = "1E293B"               # Azul oscuro para tablas
    COLOR_ZEBRA_BG = "F8FAFC"                # Gris suave alterno

    # 1. PORTADA FORMAL
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst1 = p_inst.add_run("UNIVERSIDAD SANTO TOMÁS\n")
    run_inst1.bold = True
    run_inst1.font.size = Pt(13)
    run_inst1.font.color.rgb = COLOR_SUBTITULO
    run_inst2 = p_inst.add_run("FACULTAD DE INGENIERÍA — CARRERA DE INGENIERÍA EN INFORMÁTICA\nMINERÍA DE DATOS (IEI-067)\n")
    run_inst2.font.size = Pt(11)
    run_inst2.font.color.rgb = COLOR_SUBTITULO

    p_espacio = doc.add_paragraph()
    p_espacio.paragraph_format.space_before = Pt(24)

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tit = p_tit.add_run("INFORME TÉCNICO FINAL\n")
    run_tit.bold = True
    run_tit.font.size = Pt(24)
    run_tit.font.color.rgb = COLOR_TITULO

    run_sub = p_tit.add_run("SoundData Analytics: Segmentación de Audiencias Musicales y Predicción de Éxito Comercial\n")
    run_sub.bold = True
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = COLOR_PRIMARIO

    run_desc = p_tit.add_run("Aplicación Completa de la Metodología CRISP-DM sobre un Catálogo de 85.000 Canciones de Spotify (2015–2025)\n")
    run_desc.italic = True
    run_desc.font.size = Pt(12)
    run_desc.font.color.rgb = COLOR_SUBTITULO

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_before = Pt(40)
    p_meta.paragraph_format.space_after = Pt(20)
    
    # Cuadro de Integrantes y Docentes
    tabla_portada = doc.add_table(rows=8, cols=2)
    tabla_portada.alignment = WD_TABLE_ALIGNMENT.CENTER
    datos_portada = [
        ("Docentes Guía:", "Florentino Vargas & Rosa Rao"),
        ("Grupo:", "Grupo 4"),
        ("Vicente Muñoz:", "Coordinador General, Full-Stack Lead & Integración Web"),
        ("Juan Ortiz:", "Ingeniero de Datos (Ingesta, Limpieza IQR y Preprocesamiento)"),
        ("Jordan Murillo:", "Científico de Datos (Minería de Reglas de Asociación con Apriori)"),
        ("Jorge Moncada:", "Científico de Datos (Segmentación de Canciones con K-Means)"),
        ("Jose Mendez:", "Especialista Machine Learning (Clasificación con Árbol de Decisión)"),
        ("Bastian Parraguez:", "Ingeniero de Calidad y Despliegue (QA, Pytest y CI/CD en Render)")
    ]

    for idx, (etiqueta, valor) in enumerate(datos_portada):
        c1 = tabla_portada.cell(idx, 0)
        c2 = tabla_portada.cell(idx, 1)
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(etiqueta)
        r1.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_TITULO
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(valor)
        r2.font.size = Pt(10)
        r2.font.color.rgb = COLOR_TEXTO
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_margins(c2, 60, 60, 100, 100)

    p_fecha = doc.add_paragraph()
    p_fecha.paragraph_format.space_before = Pt(40)
    p_fecha.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fec = p_fecha.add_run("Santiago de Chile — Octubre de 2026\nRepositorio Oficial: https://github.com/mkinaa/MineriaDatos")
    r_fec.font.size = Pt(10)
    r_fec.font.color.rgb = COLOR_SUBTITULO

    # Salto de página para el cuerpo del informe
    doc.add_page_break()

    # LECTURA Y PROCESAMIENTO DEL MARKDOWN
    md_path = os.path.join("docs", "INFORME_TECNICO_FINAL.md")
    with open(md_path, "r", encoding="utf-8") as f:
        lineas = f.readlines()

    # Omitir la portada repetida del Markdown (primeras ~26 líneas)
    cuerpo = []
    saltar = True
    for l in lineas:
        if l.strip().startswith("## 1. Fase 1: Comprensión del Negocio"):
            saltar = False
        if not saltar:
            cuerpo.append(l)

    en_tabla = False
    filas_tabla = []
    en_codigo = False
    lineas_codigo = []

    for linea in cuerpo:
        linea_str = linea.rstrip("\r\n")

        # Bloques de código (monospace)
        if linea_str.startswith("```"):
            if en_codigo:
                # Cerrar bloque de código
                p_code = doc.add_paragraph()
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(8)
                p_code.paragraph_format.line_spacing = 1.05
                run_code = p_code.add_run("\n".join(lineas_codigo))
                run_code.font.name = "Consolas"
                run_code.font.size = Pt(8.5)
                run_code.font.color.rgb = RGBColor(30, 41, 59)
                lineas_codigo = []
                en_codigo = False
            else:
                en_codigo = True
                lineas_codigo = []
            continue

        if en_codigo:
            lineas_codigo.append(linea_str)
            continue

        # Tablas Markdown
        if "|" in linea_str and linea_str.strip().startswith("|") and linea_str.strip().endswith("|"):
            # Ignorar filas de separación Markdown |--|--|
            if re.match(r"^\|(\s*[-:]+\s*\|)+$", linea_str.strip()):
                continue
            celdas = [c.strip() for c in linea_str.strip().split("|")[1:-1]]
            filas_tabla.append(celdas)
            en_tabla = True
            continue
        else:
            if en_tabla:
                # Renderizar tabla acumulada
                if filas_tabla:
                    num_filas = len(filas_tabla)
                    num_cols = max(len(r) for r in filas_tabla)
                    t = doc.add_table(rows=num_filas, cols=num_cols)
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    for i_f, fila in enumerate(filas_tabla):
                        es_cabecera = (i_f == 0)
                        for i_c, texto_celda in enumerate(fila):
                            celda = t.cell(i_f, i_c)
                            set_cell_margins(celda, 100, 100, 140, 140)
                            p = celda.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            
                            # Limpieza de markdown simple dentro de celda
                            texto_limpio = texto_celda.replace("**", "").replace("*", "").replace("`", "").replace("$$", "").replace("$", "")
                            
                            r = p.add_run(texto_limpio)
                            r.font.size = Pt(9)
                            if es_cabecera:
                                r.bold = True
                                r.font.color.rgb = RGBColor(255, 255, 255)
                                set_cell_background(celda, COLOR_HEADER_BG)
                            else:
                                r.font.color.rgb = COLOR_TEXTO
                                if i_f % 2 == 0:
                                    set_cell_background(celda, COLOR_ZEBRA_BG)
                    p_sep = doc.add_paragraph()
                    p_sep.paragraph_format.space_before = Pt(4)
                filas_tabla = []
                en_tabla = False

        # Títulos y Encabezados
        if linea_str.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(linea_str.replace("## ", "").strip())
            r.bold = True
            r.font.size = Pt(15)
            r.font.color.rgb = COLOR_TITULO
            continue

        if linea_str.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(linea_str.replace("### ", "").strip())
            r.bold = True
            r.font.size = Pt(12.5)
            r.font.color.rgb = COLOR_PRIMARIO
            continue

        if linea_str.startswith("#### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(linea_str.replace("#### ", "").strip())
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = COLOR_SUBTITULO
            continue

        # Listas con viñetas
        if linea_str.strip().startswith("- ") or linea_str.strip().startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            texto_vi = re.sub(r"^[\-\*]\s+", "", linea_str.strip())
            
            # Formateo inline básico (**negrita**)
            partes = re.split(r"(\*\*.*?\*\*)", texto_vi)
            for part in partes:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.size = Pt(10)
                else:
                    limpio = part.replace("`", "").replace("$$", "").replace("$", "")
                    r = p.add_run(limpio)
                    r.font.size = Pt(10)
                r.font.color.rgb = COLOR_TEXTO
            continue

        # Imágenes Markdown ![Caption](path)
        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)", linea_str.strip())
        if img_match:
            caption_text = img_match.group(1)
            raw_path = img_match.group(2)
            
            candidatos = [
                raw_path,
                os.path.join("docs", raw_path),
                os.path.join("docs", raw_path.replace("docs/", "").replace("docs\\", "")),
                raw_path.replace("docs/", "").replace("docs\\", "")
            ]
            real_img_path = next((p for p in candidatos if os.path.exists(p)), None)

            if real_img_path:
                w = Inches(6.0)
                if "matriz" in real_img_path.lower():
                    w = Inches(4.2)
                elif "arbol" in real_img_path.lower():
                    w = Inches(6.3)
                
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(10)
                p_img.paragraph_format.space_after = Pt(2)
                r_img = p_img.add_run()
                r_img.add_picture(real_img_path, width=w)

                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(2)
                p_cap.paragraph_format.space_after = Pt(12)
                r_cap = p_cap.add_run(caption_text)
                r_cap.font.size = Pt(8.5)
                r_cap.font.italic = True
                r_cap.font.color.rgb = COLOR_SUBTITULO
            continue

        # Líneas horizontales separadoras
        if linea_str.strip() in ["---", "***"]:
            continue

        # Párrafos normales
        if linea_str.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            partes = re.split(r"(\*\*.*?\*\*)", linea_str.strip())
            for part in partes:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                    r.font.size = Pt(10)
                else:
                    limpio = part.replace("`", "").replace("$$", "").replace("$", "")
                    r = p.add_run(limpio)
                    r.font.size = Pt(10)
                r.font.color.rgb = COLOR_TEXTO

    # Guardar documento
    output_docx = os.path.join("docs", "INFORME_TECNICO_FINAL.docx")
    doc.save(output_docx)
    print(f"Documento Word creado exitosamente en: {output_docx}")

if __name__ == "__main__":
    crear_documento_word()
